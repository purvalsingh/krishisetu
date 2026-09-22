"""
Checks a built deck against the SIH 2026 format rules, which are stated on the
instructions slide of the issued template:

  - at most six slides, title slide included
  - no paragraphs; points, diagrams, infographics and pictures instead
  - the provided template is used, with its section headings left alone
  - the file is submitted as PDF

Run: python3 .deck/check_sih.py <deck.pptx>
"""
import sys
from pptx import Presentation

TEMPLATE = "docs/templates/SIH2026-IDEA-Presentation-Format.pptx"

# Slides 3-6 carry fixed section headings. Slide 1 ("TITLE PAGE") and slide 2
# ("IDEA TITLE") are placeholders for our own words, not headings to preserve.
FIXED_TITLES = {
    3: "TECHNICAL APPROACH",
    4: "FEASIBILITY AND VIABILITY",
    5: "IMPACT AND BENEFITS",
    6: "RESEARCH AND REFERENCES",
}

# A run longer than this reads as a paragraph on a slide. Footnote caveats are
# allowed to be longer, so they are matched by their smaller font size instead.
PARAGRAPH_CHARS = 150
FOOTNOTE_MAX_PT = 9.5


def runs(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if not text:
                continue
            sizes = [r.font.size.pt for r in para.runs if r.font.size]
            yield text, (min(sizes) if sizes else None)


def check(path):
    prs = Presentation(path)
    problems = []

    if len(prs.slides) > 6:
        problems.append(f"{len(prs.slides)} slides; the limit is six including the title slide")

    for number, expected in FIXED_TITLES.items():
        if number > len(prs.slides):
            problems.append(f"slide {number} is missing")
            continue
        slide = prs.slides[number - 1]
        found = [t for t, _ in runs(slide) if t.upper().replace("  ", " ") == expected]
        if not found:
            problems.append(f"slide {number}: section heading '{expected}' not found")

    # Slide 2's heading is the template's "IDEA TITLE" placeholder, so it must
    # carry the name of the idea, not the body pointer repeated back.
    if len(prs.slides) > 1:
        heads = [t.upper() for t, _ in runs(prs.slides[1])]
        if any(h in ("IDEA TITLE", "PROPOSED SOLUTION") for h in heads):
            problems.append("slide 2: heading should be the idea's own title, not 'IDEA TITLE' or 'PROPOSED SOLUTION'")

    for i, slide in enumerate(prs.slides, 1):
        texts = [t for t, _ in runs(slide)]
        if i > 1:
            if not any("@SIH Idea submission" in t for t in texts):
                problems.append(f"slide {i}: template footer removed")
            # The slide number is a field inside its placeholder, so it has no
            # ordinary runs to read; check the placeholder's rendered text.
            numbered = any(
                sh.is_placeholder and sh.has_text_frame and sh.text_frame.text.strip() == str(i)
                for sh in slide.shapes
            )
            if not numbered and not any(t.strip() == str(i) for t in texts):
                problems.append(f"slide {i}: slide number removed")
            if any(t.strip() == "Your Team Name" for t in texts):
                problems.append(f"slide {i}: team name still says 'Your Team Name'")

        for text, size in runs(slide):
            if len(text) > PARAGRAPH_CHARS and (size is None or size > FOOTNOTE_MAX_PT):
                problems.append(f"slide {i}: paragraph of {len(text)} chars — '{text[:60]}...'")

        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            problems.append(f"slide {i}: no speaker notes")

    return problems


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "output/KrishiSetu_SIH2026_Illustrated_v6.pptx"
    found = check(target)
    print(f"Checking {target} against the SIH 2026 format\n")
    for p in found:
        print("  ✗", p)
    print(f"\n{len(found)} problem(s)." if found else "\nNo problems found.")
    print("Reminder: the portal accepts the PDF only.")
    sys.exit(1 if found else 0)
